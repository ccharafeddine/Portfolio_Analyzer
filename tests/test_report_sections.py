"""Reports carry the Phase 3 analytics in the Graphite style (offline run)."""

import re

import pytest

from src.reports.sections import analytics_sections


def test_sections_cover_phase3(offline_results):
    ids = [s["id"] for s in analytics_sections(offline_results)]
    for want in ("var", "drawdowns", "stress", "sectors", "factors", "dataquality"):
        assert want in ids
    for sec in analytics_sections(offline_results):
        for tb in sec["tables"]:
            assert len(tb["align"]) == len(tb["columns"])
            assert all(len(r) == len(tb["columns"]) for r in tb["rows"])


def test_html_report_graphite_and_sections(offline_results):
    from src.reports.html_builder import build_html_report

    html = build_html_report(offline_results)
    for title in ("Value at Risk", "Largest Drawdowns", "Stress Tests",
                  "Sector Attribution (Brinson-Fachler)", "Factor Exposures (Fama-French 3)",
                  "Data Quality"):
        assert f">{title}</h2>" in html, title
    assert "linear-gradient" not in html and "tabular-nums" in html
    assert not re.search(r'<h2 id="[^"]+">\d+\.', html)  # no mixed numbering
    assert 'class="r neg"' in html or 'class="r pos"' in html


def test_pdf_report_builds_with_sections(offline_results):
    from src.reports.pdf_builder import build_pdf_report

    pdf = build_pdf_report(offline_results)
    assert pdf[:4] == b"%PDF" and len(pdf) > 10_000


def test_pptx_has_analytics_slides(offline_results, tmp_path, monkeypatch):
    pytest.importorskip("pptx")
    from pptx import Presentation

    from src.reports import pptx_builder
    from src.reports.pptx_builder import build_pptx

    # Chart PNG export (kaleido + Chrome) is slow and covered elsewhere.
    monkeypatch.setattr(pptx_builder._Deck, "chart", lambda self, *a, **k: None)

    out = tmp_path / "deck.pptx"
    build_pptx(offline_results, str(out))
    titles = []
    for slide in Presentation(str(out)).slides:
        for shp in slide.shapes:
            if shp.has_text_frame and shp.text_frame.text:
                titles.append(shp.text_frame.text)
                break
    assert any(t.startswith("Value at Risk") for t in titles)
    assert any(t.startswith("Sector Attribution") for t in titles)


def test_sign_class_rules():
    from src.reports.sections import sign_class

    assert sign_class("+1.20%", True) == "pos" and sign_class("−0.5%", True) == "neg"
    assert sign_class("+0.00%", True) == "" and sign_class("−3.0%", False) == ""
